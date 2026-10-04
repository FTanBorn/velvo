# Velvo UX Scenario (LG şablonu v4.3 bölüm yapısı) -> .pptx
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
UX, OUT, ICON = sys.argv[1], sys.argv[2], sys.argv[3]
PINK, DARK, GRAY, WHITE = RGBColor(0xE6, 0x00, 0x5C), RGBColor(0x1C, 0x08, 0x12), RGBColor(0x55, 0x55, 0x5F), RGBColor(0xFF, 0xFF, 0xFF)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]

def text(slide, x, y, w, h, paras, size=14, color=DARK, bold=False, align=None, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    if anchor: tf.vertical_anchor = anchor
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        b = bold
        if isinstance(p, tuple): p, b = p
        para.text = p; para.font.size = Pt(size); para.font.color.rgb = color; para.font.bold = b; para.font.name = 'Arial'
        para.space_after = Pt(4)
        if align: para.alignment = align
    return tb

def page(title, sub=None):
    s = prs.slides.add_slide(BL)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9)); bar.fill.solid(); bar.fill.fore_color.rgb = DARK; bar.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.9), prs.slide_width, Inches(0.06)); acc.fill.solid(); acc.fill.fore_color.rgb = PINK; acc.line.fill.background()
    text(s, 0.4, 0.14, 10, 0.6, title, 24, WHITE, True)
    if sub: text(s, 9.0, 0.24, 4.0, 0.5, sub, 12, RGBColor(0xC9, 0xCD, 0xD4), align=PP_ALIGN.RIGHT)
    return s

def table(s, x, y, w, rows, head=('No.', 'Description'), widths=(0.55,), size=11):
    t = s.shapes.add_table(len(rows) + 1, len(head), Inches(x), Inches(y), Inches(w), Inches(0.3 * (len(rows) + 1))).table
    t.columns[0].width = Inches(widths[0])
    if len(head) == 2: t.columns[1].width = Inches(w - widths[0])
    for j, hd in enumerate(head):
        c = t.cell(0, j); c.text = hd; c.fill.solid(); c.fill.fore_color.rgb = PINK
        p = c.text_frame.paragraphs[0]; p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = 'Arial'
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xF6, 0xF1, 0xF4) if i % 2 else WHITE
            for p in c.text_frame.paragraphs: p.font.size = Pt(size); p.font.color.rgb = DARK; p.font.name = 'Arial'
            c.margin_top = c.margin_bottom = Emu(30000)
    return t

def shot_page(title, sub, img, rows, note=None):
    s = page(title, sub)
    s.shapes.add_picture('%s/%s.png' % (UX, img), Inches(0.4), Inches(1.25), width=Inches(7.6))
    table(s, 8.25, 1.25, 4.7, [(i + 1, r) for i, r in enumerate(rows)], size=10.5)
    if note: text(s, 0.4, 5.65, 7.6, 1.6, note, 11, GRAY)
    return s

# Kapak
s = prs.slides.add_slide(BL)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
s.shapes.add_picture(ICON, Inches(0.9), Inches(1.2), width=Inches(1.5))
text(s, 0.9, 2.9, 11, 1, 'Velvo', 54, WHITE, True)
text(s, 0.9, 3.9, 11, 0.6, 'UX Scenario Document (File Version 4.3)', 22, RGBColor(0xFF, 0x4D, 0x8D))
text(s, 0.9, 4.9, 11, 1.8, ['Submission Date:  2026 / __ / __', 'App Ver.:  1.0.0', 'App Developer:  FTanBorn', 'E-mail address:  (support e-mail)'], 16, WHITE)

# İçindekiler
s = page('Table of Contents')
text(s, 0.8, 1.4, 11, 5.5, ['1. Basic Information', '2. Document History', '3. Page Description (how to read this document)', '4. Detailed Login Information',
     '5. Main Page Description  (5-1 Start page · 5-2 Main page)', '6. Sub Page Description  (flow chart, movie, series, player, live TV, search, settings)',
     '7. Paid Content', '8. In-App Ad', 'Appendix: Notes for LG QA'], 20)

# 1. Temel bilgiler
s = page('1. Basic Information')
table(s, 0.6, 1.3, 12.1, [
  ('App title', 'Velvo'), ('App ID / version', 'io.github.ftanborn.velvo / 1.0.0'),
  ('What it is', 'A media player for the user\'s own M3U playlist or Xtream Codes account. The app contains no channels, movies or series of its own.'),
  ('Main feature', 'Choosing the audio and embedded subtitle language of the video being watched, with subtitle size, color and position settings.'),
  ('Price / payment / ads', 'Free. No in-app purchases, no payment, no advertisements.'),
  ('Accounts', 'No Velvo account. The user enters a playlist address; test playlist in section 4.'),
  ('Interface languages', 'English and Turkish. "Automatic" follows the TV language; can be changed in Settings › Language and subtitles.'),
  ('Remote controls', 'Magic Remote (pointer, OK, wheel) and the standard remote (arrows, OK, BACK, color and media keys).'),
  ('Data', 'Playlist details, profiles and preferences are stored only on the TV. Privacy policy: https://ftanborn.github.io/velvo/privacy.html'),
  ('Source code', 'https://github.com/FTanBorn/velvo (GPL-3.0)'),
], head=('Item', 'Information'), widths=(2.6,), size=12)

# 2. Belge geçmişi
s = page('2. Document History')
t = s.shapes.add_table(2, 5, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.8)).table
for j, (hd, v) in enumerate(zip(['No.', 'Sent', 'Version', 'File Name', 'Contents'], ['1', '__.__.2026', '1.0.0', 'UX Scenario', 'Initial document'])):
    for i, val in enumerate([hd, v]):
        c = t.cell(i, j); c.text = val
        p = c.text_frame.paragraphs[0]; p.font.size = Pt(13); p.font.name = 'Arial'; p.font.bold = i == 0
        p.font.color.rgb = WHITE if i == 0 else DARK
        c.fill.solid(); c.fill.fore_color.rgb = PINK if i == 0 else RGBColor(0xF6, 0xF1, 0xF4)

# 3. Sayfa açıklaması
s = page('3. Page Description')
text(s, 0.8, 1.4, 11.5, 4, ['Each page shows a screenshot of the app with numbered pink markers.', 'The table on the right describes the menu or UI button with the same number.',
     'All screenshots were taken with the English interface and the demo playlist from section 4.', 'Every button can be used with the arrow keys + OK and with the Magic Remote pointer.'], 18)

# 4. Giriş
shot_page('4. Detailed Login Information', 'Login page', 'a-login', [
  'Xtream: server address, username and password from the user\'s provider. Not needed for QA.',
  'M3U playlist: select this for QA.',
  'Playlist address: enter  https://ftanborn.github.io/velvo/demo.m3u  (LG virtual keyboard).',
  'Connect: checks the playlist and opens the Main page.'],
  ['Test Accounts (ID / PW): not required.   Activation Codes: not required.',
   'Test playlist (M3U): https://ftanborn.github.io/velvo/demo.m3u. It works on any number of devices and contains only openly licensed Blender Foundation films and public test streams.',
   'The login page appears on first launch. The playlist stays saved after relaunch and TV reboot; Settings › IPTV account › Sign out removes it.'])

# 5-1 Başlangıç
shot_page('5. Main Page Description', '5-1 Start page (profiles)', 'a-profiles', [
  'Profile: OK opens the Main page for this profile.',
  'Add profile: enter a name with the on-screen keyboard and pick a color.',
  'Manage profiles: rename, recolor or delete profiles.'],
  ['This page appears at launch when the TV has more than one profile. With a single profile the app opens the Main page directly.',
   'Each profile has its own language preferences, watch progress and My List.'])

# 5-2 Ana sayfa
shot_page('5. Main Page Description', '5-2 Main page (Home)', 'a-home', [
  'Profile: switch to another profile.', 'Search', 'Home (this page)', 'Movies: categories and titles', 'Series: categories and titles',
  'Live TV: categories and channels', 'My List: saved movies and series', 'Settings',
  'Play: plays the featured title.', 'Details: opens the detail page.', 'Add to / remove from My List',
  'Continue watching: resumes where the user stopped.', 'Content rows: one row per category; ▼ or the Magic Remote wheel scrolls down.'],
  ['The menu on the left (1–8) is reachable from every page with ◀.', 'Titles open the detail page with OK.'])

# 6-1 Akış şeması
s = page('6. Sub Page Description', '6-1 Flow chart')
def box(x, y, w, h, label, fill=RGBColor(0xF6, 0xF1, 0xF4), fc=DARK):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = PINK
    tf = b.text_frame; tf.text = label; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs: p.alignment = PP_ALIGN.CENTER; p.font.size = Pt(12); p.font.color.rgb = fc; p.font.name = 'Arial'; p.font.bold = True
    return b
def arrow(x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.line.color.rgb = GRAY; c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln(); te = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}); ln.append(te)
box(0.5, 1.5, 1.8, 0.7, 'App launch', PINK, WHITE)
box(3.0, 1.5, 2.2, 0.7, 'Login (4)\nno playlist saved')
box(3.0, 2.6, 2.2, 0.7, 'Profiles (5-1)\n2+ profiles')
box(6.0, 2.0, 2.0, 0.9, 'Main page\nHome (5-2)', PINK, WHITE)
arrow(2.3, 1.85, 3.0, 1.85); arrow(2.3, 1.85, 3.0, 2.95); arrow(5.2, 1.85, 6.0, 2.3); arrow(5.2, 2.95, 6.0, 2.6)
menu = ['Search (6-6)', 'Movies', 'Series', 'Live TV (6-5)', 'My List', 'Settings (6-7)']
for i, m in enumerate(menu):
    box(9.0, 1.3 + i * 0.62, 2.0, 0.5, m); arrow(8.0, 2.45, 9.0, 1.55 + i * 0.62)
box(6.0, 4.4, 2.0, 0.8, 'Detail\nmovie (6-2) / series (6-3)')
box(9.0, 5.3, 2.0, 0.8, 'Player (6-4)', PINK, WHITE)
arrow(7.0, 2.9, 7.0, 4.4); arrow(8.0, 4.8, 9.0, 5.6); arrow(10.0, 4.7, 10.0, 5.3)
text(s, 0.5, 4.3, 5.2, 2.6, ['With a saved playlist and one profile, the app opens the Main page directly.', 'BACK returns to the previous page.', 'BACK on the Main page exits the app (webOS.platformBack).', 'Live TV channels open the player directly; CH+ / CH− switch channels.', 'Search, Movies, Series and My List open the detail page.'], 12, GRAY)

# 6-2 Film detayı
shot_page('6. Sub Page Description', '6-2 Movie detail', 'a-movie', [
  'Language badges: audio and subtitle languages found inside the video file. The user\'s preferred languages are highlighted.',
  'Play / Continue: starts or resumes playback (shows the resume time).', 'Start over: plays from the beginning.', 'Add to / remove from My List',
  'More like this: titles from the same category.'],
  ['For MKV files Velvo reads the track list before playback ("Reading audio and subtitles…" is shown for a moment).'])

# 6-3 Dizi detayı
shot_page('6. Sub Page Description', '6-3 Series detail', 'a-series', [
  'Play: plays the next episode to watch (here S1 E1).', 'Add to / remove from My List', 'Season selector', 'Episodes: OK plays the episode; a bar shows the watched part.'],
  ['The audio and subtitle choice made in one episode is applied automatically to the next episodes of the same series.'])

# 6-4 Oynatıcı
s = page('6. Sub Page Description', '6-4 Player')
ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.25), Inches(7.6), Inches(4.275)); ph.fill.solid(); ph.fill.fore_color.rgb = RGBColor(0x11, 0x11, 0x16); ph.line.color.rgb = PINK
ph.text_frame.text = '[Player screenshot from the TV]'; ph.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x99, 0x99, 0xA0); ph.text_frame.paragraphs[0].font.size = Pt(16); ph.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
table(s, 8.25, 1.25, 4.7, [(1, 'Rewind 10 s'), (2, 'Play / Pause'), (3, 'Forward 10 s'), (4, 'Audio: list of audio tracks'),
  (5, 'Subtitles: Off + subtitle tracks, and Size, Color, Position (◀ ▶)'), (6, 'Next episode (series only)'), (7, 'Seek bar with elapsed and total time')], size=10.5)
text(s, 0.4, 5.65, 12.5, 1.7, ['Keys: any key shows the control bar · OK on ⏯ pauses/resumes · ◀ ▶ seek 10 s while the bar is hidden · YELLOW = next audio track · BLUE = next subtitle track · PLAY, PAUSE, STOP, FF, REW media keys · CH+ / CH− change channel in live TV.',
  'BACK closes an open popup first, then leaves the player and returns to the previous page.'], 11, GRAY)

# 6-5 .. 6-7
shot_page('6. Sub Page Description', '6-5 Live TV', 'a-live', ['Channel categories', 'Channels: OK plays the channel. Channels without a logo show short initials.'],
  ['In the player, CH+ / CH− switch to the previous or next channel of the same category.'])
shot_page('6. Sub Page Description', '6-6 Search', 'a-search', ['On-screen keyboard: letters and digits', 'Space', 'Delete the last character', 'Clear the search',
  'Results: movies, series and live channels whose names match'], ['OK on a result opens its detail page (or plays a channel).'])
shot_page('6. Sub Page Description', '6-7 Settings', 'a-settings', [
  'Settings tabs: Profile · Language and subtitles (app language, preferred audio and subtitle language) · Subtitle style · Home (choose and order rows) · IPTV account (sign out) · About (version, license, source code)',
  'Subtitle size', 'Subtitle color', 'Subtitle position', 'Preview of the selected style'],
  ['Subtitle style can also be changed while watching, from the Subtitles button in the player.'])

# 7, 8
s = page('7. Paid Content')
text(s, 0.8, 1.5, 11.5, 2, ['None.', 'Velvo is free. It has no in-app purchases, no payment methods and no paid content, so no test credit cards or voucher codes are needed.'], 20)
s = page('8. In-App Ad')
text(s, 0.8, 1.5, 11.5, 2, ['None.', 'Velvo shows no advertisements (no banner or video ads).'], 20)

# Ek
s = page('Appendix: Notes for LG QA')
text(s, 0.6, 1.3, 12.2, 5.8, [
  ('Content', True), 'Velvo is a player only and contains no content. Please use the demo M3U playlist from section 4; it contains only openly licensed Blender Foundation films and public test streams.',
  ('Testing subtitles and audio', True), 'Movies › "Sintel (2010)" has 1 audio track and 10 embedded subtitle tracks. In the player, open Subtitles to select a language and change size, color and position. On TV models where the platform does not allow embedded subtitle selection, Velvo hides the subtitle list and shows a message instead.',
  ('Interface language', True), 'Settings › Language and subtitles › App language: Automatic (follows the TV language), Türkçe or English.',
  ('Data', True), 'All user data stays on the TV. Settings › IPTV account › Sign out removes the saved playlist.',
  ('Contact', True), 'FTanBorn · (support e-mail) · https://github.com/FTanBorn/velvo/issues'], 13)
prs.save(OUT)
print('saved', len(prs.slides), 'slides')
